import logging
import io
from telegram import Update
from telegram.ext import ApplicationBuilder, MessageHandler, filters, ContextTypes
import openpyxl
from PIL import Image, ImageDraw, ImageFont
import os

TOKEN = "8737429803:AAFB-11SpCk8Y9NUh8XvLPuQAgUy5fnWb7w"
GROUP_ID = -1002303426539

logging.basicConfig(level=logging.INFO)

# Agent names from Oylik summa sheet
AGENT_NAMES = {}

def get_agent_names(wb):
    names = {}
    try:
        ws = wb['Oylik summa ']
        for row in ws.iter_rows(min_row=3, max_row=20, values_only=True):
            if row[9] and str(row[9]).startswith('TP-') and row[1]:
                names[row[9]] = row[1]
    except:
        pass
    return names

def format_num(val):
    if val is None or val == '' or str(val) in ('#DIV/0!', '#REF!', '#N/A'):
        return '—'
    if isinstance(val, float):
        if val == int(val):
            return f"{int(val):,}".replace(',', ' ')
        return f"{val:,.1f}".replace(',', ' ')
    if isinstance(val, int):
        return f"{val:,}".replace(',', ' ')
    return str(val)

def format_pct(val):
    if val is None or str(val) in ('#DIV/0!', '#REF!', '#N/A'):
        return '—'
    if isinstance(val, (int, float)):
        return f"{val*100:.1f}%"
    return str(val)

def make_table_image(sheet_name, agent_name, date_val, rows, itogo_row):
    # Colors
    HEADER_BG = (34, 139, 34)       # dark green
    HEADER2_BG = (144, 238, 144)    # light green
    ITOGO_BG = (255, 215, 0)        # gold
    ROW_BG1 = (255, 255, 255)
    ROW_BG2 = (240, 248, 240)
    RED = (220, 50, 50)
    BLACK = (0, 0, 0)
    WHITE = (255, 255, 255)
    DARK = (30, 30, 30)

    # Column headers & widths
    cols = ['ТМЦ', 'Кол-во\nпродажи', 'Кол-во\nпродажи', 'Факт %', 'План', 'План бл',
            'Прогноз', 'Прогноз %', 'Кам шт', 'Кам бл', '60%', '90%', '100%']
    col_widths = [160, 90, 90, 75, 95, 80, 95, 85, 90, 75, 70, 70, 70]

    total_w = sum(col_widths) + 1
    row_h = 32
    header_h = 50
    top_h = 42
    n_rows = len(rows) + 1  # +1 for itogo
    total_h = top_h + header_h + n_rows * row_h + 2

    img = Image.new('RGB', (total_w, total_h), WHITE)
    draw = ImageDraw.Draw(img)

    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 11)
        font_bold = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 11)
        font_small = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 10)
    except:
        font = ImageFont.load_default()
        font_bold = font
        font_small = font

    # Top header row
    draw.rectangle([0, 0, total_w, top_h], fill=HEADER_BG)
    date_str = str(date_val)[:10] if date_val else ''
    draw.text((8, 12), date_str, fill=WHITE, font=font_bold)
    agent_text = f"{sheet_name}  |  {agent_name}"
    draw.text((180, 12), agent_text, fill=WHITE, font=font_bold)

    # Column header row
    x = 0
    y = top_h
    draw.rectangle([0, y, total_w, y + header_h], fill=HEADER2_BG)
    for i, (col, w) in enumerate(zip(cols, col_widths)):
        draw.rectangle([x, y, x + w, y + header_h], fill=HEADER2_BG, outline=(100, 150, 100))
        lines = col.split('\n')
        line_h = 14
        start_y = y + (header_h - len(lines) * line_h) // 2
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=font_bold)
            tw = bbox[2] - bbox[0]
            draw.text((x + (w - tw) // 2, start_y), line, fill=DARK, font=font_bold)
            start_y += line_h
        x += w

    # Data rows
    y = top_h + header_h
    for idx, row in enumerate(rows):
        bg = ROW_BG1 if idx % 2 == 0 else ROW_BG2
        draw.rectangle([0, y, total_w, y + row_h], fill=bg)
        x = 0
        for i, (val, w) in enumerate(zip(row, col_widths)):
            draw.rectangle([x, y, x + w, y + row_h], fill=bg, outline=(200, 220, 200))
            # Format value
            if i == 3 or i == 7:  # Факт % or Прогноз %
                text = format_pct(val)
                # Color red if low
                color = RED if isinstance(val, (int, float)) and val < 0.6 else DARK
            elif i in (10, 11, 12):
                text = format_pct(val) if isinstance(val, (int, float)) else '—'
                color = DARK
            else:
                text = format_num(val)
                color = DARK

            bbox = draw.textbbox((0, 0), text, font=font_small)
            tw = bbox[2] - bbox[0]
            tx = x + (w - tw) // 2 if i > 0 else x + 4
            draw.text((tx, y + (row_h - 12) // 2), text, fill=color, font=font_small)
            x += w
        y += row_h

    # ИТОГО row
    draw.rectangle([0, y, total_w, y + row_h], fill=ITOGO_BG)
    x = 0
    for i, (val, w) in enumerate(zip(itogo_row, col_widths)):
        draw.rectangle([x, y, x + w, y + row_h], fill=ITOGO_BG, outline=(180, 150, 0))
        if i == 3 or i == 7:
            text = format_pct(val)
        elif i in (10, 11, 12):
            text = format_pct(val) if isinstance(val, (int, float)) else '—'
        else:
            text = format_num(val)
        bbox = draw.textbbox((0, 0), text, font=font_bold)
        tw = bbox[2] - bbox[0]
        tx = x + (w - tw) // 2 if i > 0 else x + 4
        draw.text((tx, y + (row_h - 12) // 2), text, fill=DARK, font=font_bold)
        x += w

    buf = io.BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    return buf

def parse_tp_sheet(ws):
    rows = []
    itogo = None
    date_val = None
    
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        if idx == 0:
            date_val = row[0]
            continue
        if idx == 1:
            continue  # header
        if row[0] is None:
            continue
        if str(row[0]) == 'ИТОГО':
            itogo = list(row[:13])
        else:
            rows.append(list(row[:13]))
    
    return date_val, rows, itogo

async def handle_excel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    doc = update.message.document
    if not doc or not doc.file_name.endswith('.xlsx'):
        await update.message.reply_text("Iltimos, .xlsx fayl yuboring!")
        return

    await update.message.reply_text("⏳ Fayl qayta ishlanmoqda...")

    file = await context.bot.get_file(doc.file_id)
    buf = io.BytesIO()
    await file.download_to_memory(buf)
    buf.seek(0)

    try:
        wb = openpyxl.load_workbook(buf, data_only=True)
        agent_names = get_agent_names(wb)

        tp_sheets = [s for s in wb.sheetnames if s.startswith('TP-')]

        for sheet_name in tp_sheets:
            ws = wb[sheet_name]
            agent_name = agent_names.get(sheet_name, sheet_name)
            date_val, rows, itogo = parse_tp_sheet(ws)

            if not rows or itogo is None:
                continue

            img_buf = make_table_image(sheet_name, agent_name, date_val, rows, itogo)
            caption = f"📊 {sheet_name} | {agent_name}"
            await context.bot.send_photo(
                chat_id=GROUP_ID,
                photo=img_buf,
                caption=caption
            )

        await update.message.reply_text(f"✅ {len(tp_sheets)} ta TP gruppaga yuborildi!")

    except Exception as e:
        await update.message.reply_text(f"❌ Xatolik: {e}")
        raise e

def main():
    app = ApplicationBuilder().token(TOKEN).build()
    app.add_handler(MessageHandler(filters.Document.ALL, handle_excel))
    print("Bot ishga tushdi...")
    app.run_polling()

if __name__ == '__main__':
    main()
